"""Resume and cover-letter document generation.

The layout follows a classic LaTeX resume (article class, A4, 0.4in margins):

    NAME (centred)                    headline (centred)
    phone | email | location          LinkedIn | GitHub | Portfolio
    Section Title (accent colour) ─────────────────────────────────
    **Title | Company**                                   dates (right)
      • bullet

Three outputs are produced from the same content:
* PDF  - fpdf2 with a system TrueType font (full Unicode), auto-fitted to one page;
* DOCX - python-docx, same structure with right-aligned tab stops;
* TEX  - the LaTeX source in the same style, for Overleaf / pdflatex users.

Everything stays ATS-parseable: single column, real text, standard headings.
"""
from __future__ import annotations

import copy
import os
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from fpdf import FPDF

from agents.guardrails import near_duplicate
from agents.resume_tailor import clean_text, trim_to_sentence
from config import settings
from core.logging_config import get_logger

logger = get_logger("resume_generator")

PT_TO_MM = 0.3528
MARGIN_MM = 10.16  # 0.4in, as in the LaTeX template


@dataclass(frozen=True)
class TemplateStyle:
    """Visual parameters for one resume template."""

    name: str
    display_name: str
    description: str
    font_family: str      # key into _FONT_CANDIDATES
    docx_font: str
    accent: tuple[int, int, int]
    heading_rule: bool = True


TEMPLATES: dict[str, TemplateStyle] = {
    "classic": TemplateStyle(
        name="classic",
        display_name="Classic Professional",
        description="LaTeX-style serif layout with MidnightBlue section rules. Maximum ATS compatibility.",
        font_family="serif",
        docx_font="Times New Roman",
        accent=(3, 126, 145),  # xcolor dvipsnames MidnightBlue
    ),
    "modern": TemplateStyle(
        name="modern",
        display_name="Modern Tech",
        description="Same structure in a clean sans-serif with a blue accent.",
        font_family="calibri",
        docx_font="Calibri",
        accent=(37, 99, 175),
    ),
    "minimal": TemplateStyle(
        name="minimal",
        display_name="Ultra Clean ATS",
        description="Monochrome, no rules, pure typographic hierarchy.",
        font_family="sans",
        docx_font="Arial",
        accent=(30, 30, 30),
        heading_rule=False,
    ),
}

# Regular / Bold / Italic / BoldItalic TrueType files, first match wins.
_WIN_FONTS = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts"
_FONT_CANDIDATES: dict[str, list[tuple[Path, tuple[str, str, str, str]]]] = {
    "serif": [
        (_WIN_FONTS, ("times.ttf", "timesbd.ttf", "timesi.ttf", "timesbi.ttf")),
        (Path("/usr/share/fonts/truetype/liberation"), (
            "LiberationSerif-Regular.ttf", "LiberationSerif-Bold.ttf",
            "LiberationSerif-Italic.ttf", "LiberationSerif-BoldItalic.ttf")),
        (Path("/usr/share/fonts/truetype/dejavu"), (
            "DejaVuSerif.ttf", "DejaVuSerif-Bold.ttf", "DejaVuSerif-Italic.ttf", "DejaVuSerif-BoldItalic.ttf")),
    ],
    "calibri": [
        (_WIN_FONTS, ("calibri.ttf", "calibrib.ttf", "calibrii.ttf", "calibriz.ttf")),
    ],
    "sans": [
        (_WIN_FONTS, ("arial.ttf", "arialbd.ttf", "ariali.ttf", "arialbi.ttf")),
        (Path("/usr/share/fonts/truetype/liberation"), (
            "LiberationSans-Regular.ttf", "LiberationSans-Bold.ttf",
            "LiberationSans-Italic.ttf", "LiberationSans-BoldItalic.ttf")),
        (Path("/usr/share/fonts/truetype/dejavu"), (
            "DejaVuSans.ttf", "DejaVuSans-Bold.ttf", "DejaVuSans-Oblique.ttf", "DejaVuSans-BoldOblique.ttf")),
    ],
}
_FONT_FALLBACK_CHAIN = {"serif": ["serif", "sans"], "calibri": ["calibri", "sans", "serif"], "sans": ["sans", "serif"]}
_CORE_FONT = {"serif": "times", "calibri": "helvetica", "sans": "helvetica"}

# Font sizes tried (largest first) when fitting the resume onto one page.
_BODY_SIZES = (10.5, 10.2, 10.0, 9.7, 9.4, 9.1)

_CATEGORY_TITLES = {
    "language": "Languages",
    "framework": "Frameworks & Libraries",
    "ai": "AI / ML",
    "database": "Databases",
    "cloud": "Cloud & DevOps",
    "tool": "Tools",
    "domain": "Domain Expertise",
    "leadership": "Leadership & Strategy",
    "soft": "Professional Skills",
    "general": "Core Skills",
}

MAX_ROLES_DETAILED = 4
MAX_BULLETS_PER_ROLE = 5
MAX_PROJECTS = 3
MAX_EDUCATION = 3


@dataclass
class ResumeContent:
    """Everything the renderers need."""

    name: str
    headline: str
    email: str
    phone: str
    location: str
    links: list[tuple[str, str]]          # (label, url)
    summary: str
    experiences: list[dict[str, Any]]     # {title, company, dates, bullets}
    projects: list[dict[str, Any]]        # {name, tech, bullets}
    education: list[dict[str, Any]]       # {institution, dates, degree, gpa}
    skill_groups: list[tuple[str, list[str]]]
    achievements: list[str]
    certifications: list[str] = field(default_factory=list)

    def contact_items(self) -> list[str]:
        return [p for p in (self.phone, self.email, self.location) if p]


class ResumeGenerator:
    """Renders tailored resume content to PDF, DOCX and LaTeX."""

    def __init__(self, templates_dir: str | Path | None = None) -> None:
        self.templates_dir = Path(templates_dir) if templates_dir else settings.templates_path
        self.templates_dir.mkdir(parents=True, exist_ok=True)

    # -- public API --------------------------------------------------------

    async def generate(
        self,
        output_dir: str | Path,
        job_id: str,
        user_profile: dict[str, Any],
        user_name: str,
        tailored_summary: str | None,
        tailored_bullets: dict[str, Any],
        skills_order: list[str],
        experiences: list[dict[str, Any]],
        projects: list[dict[str, Any]],
        education: list[dict[str, Any]],
        template_name: str = "classic",
        achievements: list[str] | None = None,
        selected_project_ids: list[str] | None = None,
        headline: str | None = None,
        skill_groups: dict[str, list[str]] | None = None,
        skill_categories: dict[str, str] | None = None,
        certifications: list[str] | None = None,
        pdf_engine: str = "builtin",
    ) -> tuple[str, str]:
        """Build PDF, DOCX and TEX. Returns ``(docx_path, pdf_path)``.

        The .tex sits beside them; with ``pdf_engine="latex"`` it is also compiled
        to ``resume_<id>_latex.pdf`` (skipped silently if Tectonic is missing or fails).
        """
        output_path = Path(output_dir)
        output_path.mkdir(parents=True, exist_ok=True)
        style = TEMPLATES.get(template_name, TEMPLATES["classic"])

        content = self._build_content(
            user_profile=user_profile,
            user_name=user_name,
            headline=headline,
            tailored_summary=tailored_summary,
            tailored_bullets=tailored_bullets or {},
            skills_order=skills_order or [],
            skill_groups=skill_groups,
            skill_categories=skill_categories or {},
            experiences=experiences or [],
            projects=projects or [],
            education=education or [],
            achievements=achievements or [],
            selected_project_ids=selected_project_ids or [],
            certifications=certifications or [],
        )

        # The built-in fitter trims content in place; LaTeX does its own fitting.
        latex_content = copy.deepcopy(content)
        pdf, body_size = self._fit_pdf(content, style)

        safe_id = re.sub(r"[^A-Za-z0-9_-]", "_", str(job_id))[:60]
        docx_path = output_path / f"resume_{safe_id}.docx"
        pdf_path = output_path / f"resume_{safe_id}.pdf"
        tex_path = output_path / f"resume_{safe_id}.tex"

        pdf.output(str(pdf_path))
        self._render_docx(content, style, docx_path, body_size)
        variants = latex_variants(latex_content)
        tex_path.write_text(variants[0], encoding="utf-8")

        logger.info("Generated resume %s (%s, %.1fpt, %d page(s))", pdf_path.name, style.name, body_size, pdf.page_no())

        if pdf_engine == "latex":
            from services.latex_compiler import LatexUnavailable, compile_resume, is_available

            latex_pdf = output_path / f"resume_{safe_id}_latex.pdf"
            if is_available():
                try:
                    await compile_resume(variants, latex_pdf)
                except LatexUnavailable as exc:
                    # Never serve a stale LaTeX PDF for changed content.
                    latex_pdf.unlink(missing_ok=True)
                    latex_pdf.with_suffix(".key").unlink(missing_ok=True)
                    logger.warning("LaTeX PDF unavailable, using built-in PDF: %s", exc)
        return str(docx_path), str(pdf_path)

    async def generate_cover_letter(
        self,
        output_dir: str | Path,
        job_id: str,
        user_name: str,
        user_profile: dict[str, Any],
        company: str,
        job_title: str,
        body: str,
        template_name: str = "classic",
    ) -> str:
        """Render the cover letter as a standalone PDF with the resume's header style."""
        output_path = Path(output_dir)
        output_path.mkdir(parents=True, exist_ok=True)
        style = TEMPLATES.get(template_name, TEMPLATES["classic"])

        safe_id = re.sub(r"[^A-Za-z0-9_-]", "_", str(job_id))[:60]
        pdf_path = output_path / f"cover_letter_{safe_id}.pdf"

        pdf = FPDF(format="A4", unit="mm")
        pdf.set_margins(22, 20, 22)
        pdf.set_auto_page_break(auto=True, margin=20)
        pdf.add_page()
        family, unicode_ok = _register_fonts(pdf, style.font_family)
        text = (lambda s: s) if unicode_ok else _ascii

        pdf.set_font(family, "B", 18)
        pdf.set_text_color(20, 20, 20)
        pdf.cell(0, 9, text((user_name or "Applicant").upper()), align="C", new_x="LMARGIN", new_y="NEXT")
        contact = "  |  ".join(
            p for p in (user_profile.get("phone", ""), user_profile.get("email", ""), user_profile.get("location", "")) if p
        )
        if contact:
            pdf.set_font(family, "", 10)
            pdf.set_text_color(70, 70, 70)
            pdf.cell(0, 6, text(contact), align="C", new_x="LMARGIN", new_y="NEXT")
        y = pdf.get_y() + 2
        pdf.set_draw_color(*style.accent)
        pdf.set_line_width(0.4)
        pdf.line(pdf.l_margin, y, pdf.w - pdf.r_margin, y)
        pdf.set_y(y + 7)

        pdf.set_text_color(20, 20, 20)
        pdf.set_font(family, "B", 11)
        heading = f"Application: {job_title}" + (f" at {company}" if company else "")
        pdf.multi_cell(0, 6, text(heading), new_x="LMARGIN", new_y="NEXT")
        pdf.ln(4)

        pdf.set_font(family, "", 11)
        for paragraph in (body or "").split("\n"):
            line = paragraph.strip()
            if not line:
                pdf.ln(2.5)
                continue
            pdf.multi_cell(0, 5.6, text(line), align="J", new_x="LMARGIN", new_y="NEXT")
            pdf.ln(1.5)

        pdf.output(str(pdf_path))
        logger.info("Generated cover letter %s", pdf_path.name)
        return str(pdf_path)

    # -- content assembly --------------------------------------------------

    def _build_content(
        self,
        user_profile: dict[str, Any],
        user_name: str,
        headline: str | None,
        tailored_summary: str | None,
        tailored_bullets: dict[str, Any],
        skills_order: list[str],
        skill_groups: dict[str, list[str]] | None,
        skill_categories: dict[str, str],
        experiences: list[dict[str, Any]],
        projects: list[dict[str, Any]],
        education: list[dict[str, Any]],
        achievements: list[str],
        selected_project_ids: list[str],
        certifications: list[str],
    ) -> ResumeContent:
        # Experience: roles chosen by the tailor get its bullets; if nothing was
        # tailored, the most recent roles keep their own bullets; the rest are
        # listed on one line each, so the full career history stays visible.
        tailored_ids = {str(k) for k, v in tailored_bullets.items() if v}
        formatted_experiences: list[dict[str, Any]] = []
        shown: list[str] = []  # every bullet already on the page, to avoid repeats
        detailed = 0
        for index, exp in enumerate(experiences[:14]):
            exp_id = str(exp.get("id", index))
            own_from = None
            if exp_id in tailored_ids:
                # Tailored lines first, then the role's own lines (the user's words) to fill
                # the role; lines covering the same thing are skipped, page fitting trims extras.
                candidates = list(tailored_bullets[exp_id]) + list(exp.get("bullets") or [])
                own_from = len(tailored_bullets[exp_id])
                limit = MAX_BULLETS_PER_ROLE
            elif detailed < MAX_ROLES_DETAILED + 2:
                # Roles the tailor did not pick still show their own top lines;
                # the page-fitting step trims these first if space runs out.
                candidates = exp.get("bullets") or []
                limit = 2 if tailored_ids else MAX_BULLETS_PER_ROLE
            else:
                candidates, limit = [], 0
            bullets = []
            for position, item in enumerate(candidates):
                text = _clean_bullet(item)
                if own_from is not None and position >= own_from and any(_same_point(text, b) for b in bullets):
                    continue  # the tailored version of this line is already there
                if text and not any(near_duplicate(text, prior) for prior in shown):
                    bullets.append(text)
                    shown.append(text)
                if len(bullets) >= limit:
                    break
            if bullets:
                detailed += 1
            title = str(exp.get("title") or "").strip()
            company = str(exp.get("company") or "").strip()
            formatted_experiences.append({
                "title": " | ".join(p for p in (title, company) if p),
                "dates": _format_dates(exp.get("start_date"), exp.get("end_date"), exp.get("is_current")),
                "bullets": bullets,
                "tailored": exp_id in tailored_ids,
            })

        # The tailor's picks first, then the user's other projects, so a short tailored
        # selection never leaves the page half empty (page fitting trims extras first).
        picked = set(selected_project_ids)
        chosen = [p for p in projects if str(p.get("id")) in picked] + [p for p in projects if str(p.get("id")) not in picked]
        formatted_projects = []
        for project in chosen[:MAX_PROJECTS]:
            if not project.get("name"):
                continue
            tech = [str(t) for t in (project.get("technologies") or [])][:5]
            tech = [t for t in tech if t.lower() not in _NOT_TECH]
            sentences = re.split(r"(?<=[.!?])\s+", clean_text(project.get("description") or ""))
            formatted_projects.append({
                "name": " | ".join(p for p in (str(project["name"]).strip(), ", ".join(tech)) if p),
                "bullets": [_clean_bullet(trim_to_sentence(s, 200)) for s in sentences if len(s) > 15][:3],
            })
        if sum(1 for p in formatted_projects if p["bullets"]) >= 2:
            # A bare repo name says nothing to a recruiter once real projects are listed.
            formatted_projects = [p for p in formatted_projects if p["bullets"]]

        formatted_education = []
        for edu in education[:MAX_EDUCATION + 2]:
            if not (edu.get("institution") or edu.get("degree")):
                continue
            degree = str(edu.get("degree") or "").strip()
            field_name = str(edu.get("field") or "").strip()
            gpa = str(edu.get("gpa") or "").strip()
            institution = str(edu.get("institution") or "").strip()
            title = f"{degree} in {field_name}" if degree and field_name else (degree or field_name)
            same = next((e for e in formatted_education if institution and e["institution"].lower() == institution.lower()), None)
            if same is not None:
                # Diplomas / certificates from the same school: one line under the main degree.
                year = _format_dates(edu.get("start_date"), edu.get("end_date"), False)
                same["extras"].append(f"{title} ({year})" if year else title)
                continue
            formatted_education.append({
                "institution": institution,
                "dates": _format_dates(edu.get("start_date"), edu.get("end_date"), False),
                "degree": title,
                "gpa": f"CGPA: {gpa}" if gpa and not gpa.lower().startswith(("cgpa", "gpa")) else gpa,
                "extras": [],
            })
        for entry in formatted_education:
            if entry["extras"]:
                entry["degree"] = f"{entry['degree']}; also {', '.join(entry['extras'])}"
        formatted_education = formatted_education[:MAX_EDUCATION]

        groups = self._skill_groups(skills_order, skill_groups, skill_categories)

        links = []
        for label, key in (("LinkedIn", "linkedin_url"), ("GitHub", "github_url"), ("Portfolio", "portfolio_url")):
            url = str(user_profile.get(key) or "").strip()
            if url:
                links.append((label, url if url.startswith("http") else f"https://{url}"))

        summary = tailored_summary or user_profile.get("professional_summary") or ""
        return ResumeContent(
            name=(user_name or "Applicant").strip(),
            headline=clean_text(headline or user_profile.get("headline") or ""),
            email=str(user_profile.get("email") or "").strip(),
            phone=str(user_profile.get("phone") or "").strip(),
            location=str(user_profile.get("location") or "").strip(),
            links=links,
            summary=trim_to_sentence(summary, 700),
            experiences=formatted_experiences,
            projects=formatted_projects,
            education=formatted_education,
            skill_groups=groups,
            achievements=_real_achievements(achievements, formatted_experiences, formatted_projects),
            certifications=[str(c).strip() for c in certifications if str(c).strip()][:5],
        )

    @staticmethod
    def _skill_groups(
        skills_order: list[str],
        skill_groups: dict[str, list[str]] | None,
        skill_categories: dict[str, str],
    ) -> list[tuple[str, list[str]]]:
        if skill_groups:
            groups = [(str(name), [str(s) for s in members]) for name, members in skill_groups.items() if members]
        else:
            buckets: dict[str, list[str]] = {}
            for skill in skills_order:
                category = (skill_categories.get(skill) or "general").lower()
                title = _CATEGORY_TITLES.get(category, category.replace("-", " ").title())
                buckets.setdefault(title, []).append(skill)
            groups = list(buckets.items())
        return _tidy_skill_groups(groups)

    # -- PDF -----------------------------------------------------------------

    def _fit_pdf(self, content: ResumeContent, style: TemplateStyle) -> tuple[FPDF, float]:
        """Largest font size that fits one page; trim gently, then allow a second page."""
        for size in _BODY_SIZES:
            pdf = self._render_pdf(content, style, size)
            if pdf.page_no() == 1:
                return pdf, size

        # Still too long at the smallest size: trim the least important content.
        for trim in (_trim_untailored, _trim_older_bullets, _trim_extras, _trim_bullets_to_three):
            trim(content)
            pdf = self._render_pdf(content, style, _BODY_SIZES[-1])
            if pdf.page_no() == 1:
                return pdf, _BODY_SIZES[-1]
        # A long senior career may need two pages - better than cutting real roles.
        return self._render_pdf(content, style, _BODY_SIZES[2]), _BODY_SIZES[2]

    def _render_pdf(self, content: ResumeContent, style: TemplateStyle, size: float) -> FPDF:
        pdf = FPDF(format="A4", unit="mm")
        pdf.set_margins(MARGIN_MM, MARGIN_MM, MARGIN_MM)
        pdf.set_auto_page_break(auto=True, margin=MARGIN_MM)
        pdf.add_page()
        family, unicode_ok = _register_fonts(pdf, style.font_family)
        writer = _PdfWriter(pdf, family, unicode_ok, style, size)
        writer.render(content)
        return pdf

    # -- DOCX ----------------------------------------------------------------

    def _render_docx(self, content: ResumeContent, style: TemplateStyle, path: Path, body_size: float) -> None:
        document = Document()
        section = document.sections[0]
        section.page_width, section.page_height = Inches(8.27), Inches(11.69)  # A4
        for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
            setattr(section, side, Inches(0.4))
        usable = section.page_width - section.left_margin - section.right_margin

        normal = document.styles["Normal"]
        normal.font.name = style.docx_font
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), style.docx_font)
        normal.font.size = Pt(body_size)
        normal.paragraph_format.space_after = Pt(0)
        normal.paragraph_format.space_before = Pt(0)
        normal.paragraph_format.line_spacing = 1.08
        accent = RGBColor(*style.accent)

        def centered(text: str, size: float, bold: bool = False, space_after: float = 1) -> None:
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(text)
            run.bold = bold
            run.font.size = Pt(size)
            paragraph.paragraph_format.space_after = Pt(space_after)

        centered(content.name.upper(), body_size + 9, bold=True, space_after=2)
        if content.headline:
            centered(content.headline, body_size + 2, space_after=3)
        if content.contact_items():
            centered("  |  ".join(content.contact_items()), body_size)
        if content.links:
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for index, (label, url) in enumerate(content.links):
                if index:
                    paragraph.add_run("  |  ")
                _add_hyperlink(paragraph, _strip_scheme(url), url, style.accent)

        def heading(title: str) -> None:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(7)
            paragraph.paragraph_format.space_after = Pt(3)
            run = paragraph.add_run(title)
            run.bold = True
            run.font.size = Pt(body_size + 2)
            run.font.color.rgb = accent
            if style.heading_rule:
                _add_bottom_border(paragraph, style.accent)

        def row(left: str, right: str = "", bold: bool = True, space_before: float = 3) -> None:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(space_before)
            paragraph.paragraph_format.tab_stops.add_tab_stop(usable, WD_TAB_ALIGNMENT.RIGHT)
            left_run = paragraph.add_run(left)
            left_run.bold = bold
            if right:
                paragraph.add_run("\t" + right)

        def bullet(text: str) -> None:
            paragraph = document.add_paragraph(text, style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(0)

        if content.summary:
            heading("Professional Summary")
            paragraph = document.add_paragraph(content.summary)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        if content.experiences:
            heading("Experience")
            for exp in content.experiences:
                row(exp["title"], exp["dates"], space_before=4 if exp["bullets"] else 2)
                for item in exp["bullets"]:
                    bullet(item)

        if content.projects:
            heading("Projects")
            for project in content.projects:
                row(project["name"], space_before=4)
                for item in project["bullets"]:
                    bullet(item)

        if content.education:
            heading("Education")
            for edu in content.education:
                row(edu["institution"], edu["dates"])
                if edu["degree"] or edu["gpa"]:
                    row(edu["degree"], edu["gpa"], bold=False, space_before=0)

        if content.skill_groups:
            heading("Technical Skills")
            for group, members in content.skill_groups:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.space_after = Pt(1)
                paragraph.add_run(f"{group}: ").bold = True
                paragraph.add_run(", ".join(members))

        if content.achievements:
            heading("Achievements")
            for item in content.achievements:
                bullet(item)

        if content.certifications:
            heading("Certifications")
            for item in content.certifications:
                bullet(item)

        document.save(str(path))


class _PdfWriter:
    """Draws one resume onto an FPDF page in the LaTeX-template style."""

    def __init__(self, pdf: FPDF, family: str, unicode_ok: bool, style: TemplateStyle, size: float) -> None:
        self.pdf = pdf
        self.family = family
        self.style = style
        self.size = size
        self.line = size * PT_TO_MM * 1.28
        self.text = (lambda s: s) if unicode_ok else _ascii
        self.bullet_glyph = "•" if unicode_ok else "-"
        self.width = pdf.w - pdf.l_margin - pdf.r_margin

    def font(self, weight: str = "", delta: float = 0.0) -> None:
        self.pdf.set_font(self.family, weight, self.size + delta)

    def render(self, c: ResumeContent) -> None:
        pdf = self.pdf
        pdf.set_text_color(15, 15, 15)

        self.font("B", 9)
        pdf.cell(0, (self.size + 9) * PT_TO_MM * 1.15, self.text(c.name.upper()), align="C", new_x="LMARGIN", new_y="NEXT")
        if c.headline:
            pdf.ln(0.8)
            self.font("", 2)
            pdf.multi_cell(0, (self.size + 2) * PT_TO_MM * 1.3, self.text(c.headline), align="C", new_x="LMARGIN", new_y="NEXT")
        pdf.ln(0.8)
        self.font()
        if c.contact_items():
            pdf.cell(0, self.line, self.text("   |   ".join(c.contact_items())), align="C", new_x="LMARGIN", new_y="NEXT")
        if c.links:
            self._links_line(c.links)

        if c.summary:
            self.heading("Professional Summary")
            self.font()
            pdf.multi_cell(0, self.line, self.text(c.summary), align="J", new_x="LMARGIN", new_y="NEXT")

        if c.experiences:
            self.heading("Experience")
            for exp in c.experiences:
                self.row(exp["title"], exp["dates"], gap=1.6 if exp["bullets"] else 0.7)
                for item in exp["bullets"]:
                    self.bullet(item)

        if c.projects:
            self.heading("Projects")
            for project in c.projects:
                self.row(project["name"], "", gap=1.6)
                for item in project["bullets"]:
                    self.bullet(item)

        if c.education:
            self.heading("Education")
            for edu in c.education:
                self.row(edu["institution"], edu["dates"], gap=1.2)
                if edu["degree"] or edu["gpa"]:
                    self.row(edu["degree"], edu["gpa"], bold=False, gap=0)

        if c.skill_groups:
            self.heading("Technical Skills")
            for group, members in c.skill_groups:
                self.font()
                pdf.multi_cell(
                    0, self.line, self.text(f"**{_md_safe(group)}:** {_md_safe(', '.join(members))}"),
                    markdown=True, new_x="LMARGIN", new_y="NEXT",
                )
                pdf.ln(0.4)

        if c.achievements:
            self.heading("Achievements")
            for item in c.achievements:
                self.bullet(item)

        if c.certifications:
            self.heading("Certifications")
            for item in c.certifications:
                self.bullet(item)

    def _links_line(self, links: list[tuple[str, str]]) -> None:
        pdf = self.pdf
        self.font()
        separator = "   |   "
        labels = [self.text(_strip_scheme(url)) for _label, url in links]
        total = sum(pdf.get_string_width(l) for l in labels) + pdf.get_string_width(separator) * (len(labels) - 1)
        pdf.set_x(pdf.l_margin + max(0.0, (self.width - total) / 2))
        for index, ((_label, url), text) in enumerate(zip(links, labels)):
            if index:
                pdf.set_text_color(15, 15, 15)
                pdf.cell(pdf.get_string_width(separator), self.line, separator)
            pdf.set_text_color(*self.style.accent)
            pdf.cell(pdf.get_string_width(text), self.line, text, link=url)
        pdf.set_text_color(15, 15, 15)
        pdf.ln(self.line)

    def heading(self, title: str) -> None:
        pdf = self.pdf
        pdf.ln(2.6)
        self.font("B", 2)
        pdf.set_text_color(*self.style.accent)
        pdf.cell(0, (self.size + 2) * PT_TO_MM * 1.25, self.text(title), new_x="LMARGIN", new_y="NEXT")
        if self.style.heading_rule:
            y = pdf.get_y() + 0.2
            pdf.set_draw_color(*self.style.accent)
            pdf.set_line_width(0.3)
            pdf.line(pdf.l_margin, y, pdf.w - pdf.r_margin, y)
        pdf.ln(1.4)
        pdf.set_text_color(15, 15, 15)

    def row(self, left: str, right: str, bold: bool = True, gap: float = 1.2) -> None:
        """Left text (bold) with right-aligned text on the same baseline, LaTeX \\hfill style."""
        pdf = self.pdf
        pdf.ln(gap)
        if pdf.get_y() + self.line > pdf.page_break_trigger:
            pdf.add_page()
        self.font()
        right_text = self.text(right or "")
        right_width = pdf.get_string_width(right_text) + 1 if right_text else 0
        y = pdf.get_y()
        if right_text:
            pdf.set_xy(pdf.l_margin + self.width - right_width, y)
            pdf.cell(right_width, self.line, right_text, align="R")
            pdf.set_xy(pdf.l_margin, y)
        self.font("B" if bold else "")
        pdf.multi_cell(self.width - right_width - 2, self.line, self.text(left), new_x="LMARGIN", new_y="NEXT")

    def bullet(self, text: str) -> None:
        pdf = self.pdf
        self.font()
        indent = 3.2
        pdf.set_x(pdf.l_margin + 1.2)
        pdf.cell(indent, self.line, self.bullet_glyph)
        pdf.multi_cell(self.width - indent - 1.2, self.line, self.text(text), new_x="LMARGIN", new_y="NEXT")


# ---------------------------------------------------------------------------
# LaTeX
# ---------------------------------------------------------------------------

_TEX_PREAMBLE = r"""\documentclass[10pt,a4paper]{article}

\usepackage[a4paper,margin=0.4in]{geometry}
\usepackage{titlesec}
\usepackage{enumitem}
\usepackage[hidelinks]{hyperref}
\usepackage[dvipsnames]{xcolor}
\usepackage{fontawesome5}
\usepackage{parskip}

\pagestyle{empty}
\setlength{\parindent}{0pt}
\setlist[itemize]{leftmargin=*,nosep}

\titleformat{\section}
{\large\bfseries\color{MidnightBlue}}
{}{0em}{}
[\color{MidnightBlue}\titlerule]

% Left text with right-aligned dates; if they don't fit, the dates move to
% the next line, still flush right (TeXbook "\hfill\penalty" idiom).
\newcommand{\resumeentry}[2]{{\parfillskip=0pt #1\nobreak\hfill\penalty50\hskip1em\null\nobreak\hfill\mbox{#2}\par}}
"""

# Progressively tighter spacing, tried in order until the resume fits one page.
# Level 0 is the template exactly as designed.
_TEX_DENSITY = [
    "",
    "\n\\setlength{\\parskip}{3pt}\n\\titlespacing*{\\section}{0pt}{8pt}{4pt}\n",
    "\n\\setlength{\\parskip}{2pt}\n\\titlespacing*{\\section}{0pt}{6pt}{3pt}\n\\linespread{0.96}\n",
    "\n\\setlength{\\parskip}{1.5pt}\n\\titlespacing*{\\section}{0pt}{5pt}{2pt}\n\\linespread{0.94}\n"
    "\\AtBeginDocument{\\small}\n",
]

_TEX_SPECIALS = {
    "\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "$": r"\$", "#": r"\#", "_": r"\_",
    "{": r"\{", "}": r"\}", "~": r"\textasciitilde{}", "^": r"\textasciicircum{}",
    # Without T1 font encoding pdflatex prints a bare "|" as an em dash.
    "|": r"\textbar{}", "<": r"\textless{}", ">": r"\textgreater{}",
}
_TEX_UNICODE = {"→": r"$\rightarrow$", "–": "--", "—": "---", "₹": "Rs.", "•": r"\textbullet{}", "…": r"\ldots{}",
                "≥": r"$\geq$", "≤": r"$\leq$", "×": r"$\times$"}


def tex_escape(text: str) -> str:
    out = "".join(_TEX_SPECIALS.get(ch, ch) for ch in str(text or ""))
    # Let long slash-joined runs ("CSV/XLSX/JSON/...") break instead of running into the margin.
    out = out.replace("/", r"/\allowbreak{}")
    for source, target in _TEX_UNICODE.items():
        out = out.replace(source, target)
    return out


def latex_variants(content: ResumeContent) -> list[str]:
    """LaTeX sources from loosest to tightest; the last one also trims older bullets."""
    variants = [render_tex(content, density) for density in range(len(_TEX_DENSITY))]
    trimmed = copy.deepcopy(content)
    _trim_untailored(trimmed)
    _trim_older_bullets(trimmed)
    variants.append(render_tex(trimmed, len(_TEX_DENSITY) - 1))
    return variants


def render_tex(c: ResumeContent, density: int = 0) -> str:
    """LaTeX source matching the user's template (compile with Tectonic, pdflatex or Overleaf)."""
    e = tex_escape
    preamble = _TEX_PREAMBLE + _TEX_DENSITY[density] + "\n" + r"\begin{document}" + "\n"
    lines = [preamble, r"\begin{center}", rf"{{\LARGE\textbf{{{e(c.name.upper())}}}}}\\[4pt]"]
    if c.headline:
        lines.append(rf"{{\large {e(c.headline)}}}\\[4pt]")
    contact = []
    if c.phone:
        contact.append(rf"\faPhone\ {e(c.phone)}")
    if c.email:
        contact.append(rf"\faEnvelope\ \href{{mailto:{c.email}}}{{{e(c.email)}}}")
    if c.location:
        contact.append(rf"\faMapMarker*\ {e(c.location)}")
    if contact:
        lines += [" \\quad\n".join(contact), ""]
    icons = {"LinkedIn": r"\faLinkedin", "GitHub": r"\faGithub", "Portfolio": r"\faGlobe"}
    if c.links:
        lines.append("\n\\quad|\\quad\n".join(rf"{icons.get(label, '')}\ \href{{{url}}}{{{label}}}" for label, url in c.links))
    lines += [r"\end{center}", ""]

    def itemize(items: list[str]) -> None:
        lines.append(r"\begin{itemize}")
        lines.extend(rf"\item {e(item)}" for item in items)
        lines.extend([r"\end{itemize}", ""])

    if c.summary:
        lines += [r"\section{Professional Summary}", "", e(c.summary), ""]
    if c.experiences:
        lines += [r"\section{Experience}", ""]
        for exp in c.experiences:
            lines += [rf"\resumeentry{{\textbf{{{e(exp['title'])}}}}}{{{e(exp['dates'])}}}", ""]
            if exp["bullets"]:
                itemize(exp["bullets"])
    if c.projects:
        lines += [r"\section{Projects}", ""]
        for project in c.projects:
            lines += [rf"\textbf{{{e(project['name'])}}}", ""]
            if project["bullets"]:
                itemize(project["bullets"])
    if c.education:
        lines += [r"\section{Education}", ""]
        for edu in c.education:
            lines += [rf"\resumeentry{{\textbf{{{e(edu['institution'])}}}}}{{{e(edu['dates'])}}}", ""]
            if edu["degree"] or edu["gpa"]:
                lines += [rf"{e(edu['degree'])} \hfill {e(edu['gpa'])}", ""]
    if c.skill_groups:
        lines += [r"\section{Technical Skills}", ""]
        for group, members in c.skill_groups:
            lines += [rf"\textbf{{{e(group)}:}} {e(', '.join(members))}", ""]
    if c.achievements:
        lines += [r"\section{Achievements}", ""]
        itemize(c.achievements)
    if c.certifications:
        lines += [r"\section{Certifications}", ""]
        itemize(c.certifications)
    lines.append(r"\end{document}")
    return "\n".join(lines) + "\n"


# ---------------------------------------------------------------------------
# Fitting helpers
# ---------------------------------------------------------------------------


def _trim_untailored(content: ResumeContent) -> None:
    """Roles the tailor did not choose drop to one line of their own text."""
    if any(exp.get("tailored") for exp in content.experiences):
        for exp in content.experiences:
            if not exp.get("tailored"):
                exp["bullets"] = exp["bullets"][:1]


def _trim_older_bullets(content: ResumeContent) -> None:
    """Keep full bullets for the first two detailed roles, two for the rest."""
    seen = 0
    for exp in content.experiences:
        if exp["bullets"]:
            seen += 1
            if seen > 2:
                exp["bullets"] = exp["bullets"][:2]


def _trim_extras(content: ResumeContent) -> None:
    content.certifications = []
    content.achievements = content.achievements[:3]
    content.projects = content.projects[:2]
    for project in content.projects:
        project["bullets"] = project["bullets"][:2]


def _trim_bullets_to_three(content: ResumeContent) -> None:
    for exp in content.experiences:
        exp["bullets"] = exp["bullets"][:3]


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _register_fonts(pdf: FPDF, family_key: str) -> tuple[str, bool]:
    """Register a Unicode TrueType family; fall back to a Latin-1 core font."""
    for key in _FONT_FALLBACK_CHAIN.get(family_key, [family_key]):
        for directory, files in _FONT_CANDIDATES.get(key, []):
            paths = [directory / name for name in files]
            if all(p.exists() for p in paths):
                name = f"resume-{key}"
                if name not in pdf.fonts and f"{name}b" not in pdf.fonts:
                    for path, weight in zip(paths, ("", "B", "I", "BI")):
                        pdf.add_font(name, weight, str(path))
                return name, True
    return _CORE_FONT.get(family_key, "helvetica"), False


def _add_bottom_border(paragraph, color: tuple[int, int, int]) -> None:
    """python-docx exposes no border API, so emit the OOXML directly."""
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "%02X%02X%02X" % color)
    borders.append(bottom)
    p_pr.append(borders)


def _add_hyperlink(paragraph, text: str, url: str, color: tuple[int, int, int]) -> None:
    part = paragraph.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color_el = OxmlElement("w:color")
    color_el.set(qn("w:val"), "%02X%02X%02X" % color)
    r_pr.append(color_el)
    run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


_LIGATURES = {
    "‘": "'", "’": "'", "“": '"', "”": '"',
    "–": "-", "—": "-", "…": "...", "•": "-", "→": "->",
    " ": " ", "‑": "-", "−": "-", "­": "",
}


def _ascii(text: str) -> str:
    """fpdf2's core fonts are Latin-1 only; normalise before encoding."""
    result = str(text or "")
    for source, target in _LIGATURES.items():
        result = result.replace(source, target)
    return result.encode("latin-1", "replace").decode("latin-1")


def _md_safe(text: str) -> str:
    """Neutralise fpdf2 markdown markers inside user text."""
    return str(text).replace("**", "* *").replace("__", "_ _").replace("--", "- -")


def _same_point(a: str, b: str) -> bool:
    """Two lines make the same point: most of the shorter one's content words appear in the other."""
    words = lambda s: {w for w in re.findall(r"[a-z0-9+]+", s.lower()) if len(w) > 3}
    wa, wb = words(a), words(b)
    if not wa or not wb:
        return False
    return len(wa & wb) / min(len(wa), len(wb)) >= 0.5


def _real_achievements(items: list[Any], experiences: list[dict], projects: list[dict]) -> list[str]:
    """Achievements that read as achievements: whole statements, not stray numbers
    ("110+ commits") or facts already stated in an experience or project line."""
    already = " ".join(b for block in experiences + projects for b in block.get("bullets", [])).lower()
    kept: list[str] = []
    for item in items:
        text = _clean_bullet(item)
        if len(text.split()) < 6:
            continue
        numbers = re.findall(r"\d[\d,.]*\+?", text)
        if numbers and all(n.lower() in already for n in numbers):
            continue
        kept.append(text)
    return kept[:5]


def _clean_bullet(text: Any) -> str:
    text = clean_text(text).lstrip("-*•· ").strip()
    # Descriptions cut off mid-sentence ("..., generates code to compute results,") end cleanly.
    if text and text[-1] in ",;:":
        text = text.rstrip(",;: ") + "."
    return text


# GitHub "languages" that are file formats, not skills.
_NOT_TECH = {"jupyter notebook", "html", "css", "shell", "dockerfile", "makefile", "procfile"}
# Internal category names that must never be printed as a heading.
_INTERNAL_GROUPS = {"auto-imported", "auto imported", "general", "other", "misc", "miscellaneous", ""}


def _tidy_skill_groups(groups: list[tuple[str, list[str]]]) -> list[tuple[str, list[str]]]:
    """Fold internal labels and one-item groups into an "Other" line; drop repeats and file formats."""
    seen: set[str] = set()
    kept: list[tuple[str, list[str]]] = []
    spill: list[str] = []
    for name, members in groups:
        unique = []
        for skill in members:
            skill = clean_text(skill)
            key = skill.lower()
            if skill and key not in seen and key not in _NOT_TECH:
                seen.add(key)
                unique.append(skill)
        if not unique:
            continue
        if name.strip().lower() in _INTERNAL_GROUPS or len(unique) < 2:
            spill.extend(unique)
        else:
            kept.append((name, unique))
    if spill:
        if kept and len(spill) < 2:
            kept[-1] = (kept[-1][0], kept[-1][1] + spill)
        else:
            kept.append(("Other" if kept else "Skills", spill))
    return kept


def _strip_scheme(url: str) -> str:
    return re.sub(r"^https?://(www\.)?", "", url or "").rstrip("/")


def _format_dates(start: Any, end: Any, is_current: Any) -> str:
    start_text = str(start or "").strip()
    end_text = "Present" if is_current else str(end or "").strip()
    if start_text and end_text and start_text != end_text:
        return f"{start_text} – {end_text}"
    return start_text or end_text or ""


def list_templates() -> list[dict[str, str]]:
    return [
        {
            "name": style.name,
            "display_name": style.display_name,
            "description": style.description,
            "font": style.docx_font,
            "accent_color": "#%02X%02X%02X" % style.accent,
        }
        for style in TEMPLATES.values()
    ]
